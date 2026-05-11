import sys
sys.stdout.reconfigure(encoding='utf-8')
from docx import Document
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from copy import deepcopy

doc = Document(r'D:\Spinrise\Docs\UI_Blueprint\SPINRISE_UIDesign_Blueprint_v1.1_FINAL_08May2026.docx')

# --- HELPERS ---
def make_run(text, bold=False):
    r = OxmlElement('w:r')
    if bold:
        rPr = OxmlElement('w:rPr')
        b = OxmlElement('w:b')
        rPr.append(b)
        r.append(rPr)
    t = OxmlElement('w:t')
    t.text = text
    t.set('{http://www.w3.org/XML/1998/namespace}space', 'preserve')
    r.append(t)
    return r

def make_para(text='', bold=False):
    p = OxmlElement('w:p')
    if text:
        p.append(make_run(text, bold))
    return p

def set_cell_content(cell, items):
    tc = cell._tc
    existing = tc.findall(qn('w:p'))
    first_pPr = None
    if existing:
        first_pPr = existing[0].find(qn('w:pPr'))
    for p in existing:
        tc.remove(p)
    for i, (text, bold) in enumerate(items):
        p = OxmlElement('w:p')
        if i == 0 and first_pPr is not None:
            p.append(deepcopy(first_pPr))
        if text:
            p.append(make_run(text, bold))
        tc.append(p)

def clear_para(para):
    p = para._element
    for child in list(p):
        if child.tag != qn('w:pPr'):
            p.remove(child)

def set_para(para, text, bold=False):
    clear_para(para)
    para._element.append(make_run(text, bold))

def insert_after(ref_elem, items):
    last = ref_elem
    for text, bold in items:
        p = make_para(text, bold)
        last.addnext(p)
        last = p

# ============================================================
# FIX 1: Section 7.1 blue box (Table[27])
# ============================================================
cell_71 = doc.tables[27].rows[0].cells[0]
set_cell_content(cell_71, [
    ('STATUS: PROPOSED — Pending CEO Decision (Fri 8 May 2026)', True),
    ('', False),
    ('Note: The comparison table and grid design standards in Section 7.3 that do not depend on library selection (row height, column alignment, field types, etc.) are CONFIRMED. The library selection itself (AG Grid Enterprise vs alternatives) is PENDING CEO approval.', False),
])
print('FIX 1 done:', doc.tables[27].rows[0].cells[0].text[:80])

# ============================================================
# FIX 2: Delete duplicate Value row in Table[17] (Row[3] = shorter one)
# ============================================================
tbl17 = doc.tables[17]
print('Before delete - Row3:', ' | '.join(c.text.strip()[:30] for c in tbl17.rows[3].cells))
tbl17._tbl.remove(tbl17.rows[3]._tr)
print('FIX 2 done. New Row3:', ' | '.join(c.text.strip()[:30] for c in tbl17.rows[3].cells))

# ============================================================
# FIX 3: Remove Budget Category para from Section 5.1.1
# ============================================================
para_to_del = None
in_511 = False
for p in doc.paragraphs:
    if '5.1.1' in p.text and 'Add Mode' in p.text:
        in_511 = True
        continue
    if in_511:
        if 'REFERENCE IMPLEMENTATION' in p.text and 'Budget Category' in p.text:
            para_to_del = p
            break
        if 'CEO Decision' in p.text:
            break

if para_to_del:
    para_to_del._element.getparent().remove(para_to_del._element)
    print('FIX 3 done: Budget Category paragraph removed from 5.1.1')
else:
    print('FIX 3 WARNING: para not found')

# ============================================================
# FIX 4: Add Architecture Pattern to Section 5.4.1
# ============================================================
para_541 = None
for p in doc.paragraphs:
    if '5.4.1' in p.text and 'Multi-Vertical' in p.text:
        para_541 = p
        break

if para_541:
    insert_after(para_541._element, [
        ('', False),
        ('Architecture Pattern:', False),
        ('This conditional display logic establishes the reference pattern for all future vertical-specific fields in SpinRise. All such fields must follow this structure:', False),
        ('  • Visibility rule explicitly documented in Blueprint', False),
        ('  • Backend configuration table maps field to division + transaction type', False),
        ('  • Frontend renders conditionally based on logged-in user’s division profile', False),
        ('', False),
        ('NOTE: Sasi to confirm backend configuration table design before this section is finalised and Blueprint is submitted for CEO sign-off.', False),
    ])
    print('FIX 4 done: Architecture Pattern added to 5.4.1')
else:
    print('FIX 4 WARNING: Para 5.4.1 not found')

# ============================================================
# FIX 5: Expand Add Multiple row in Table[18]
# ============================================================
tbl18 = doc.tables[18]
am_row = None
for row in tbl18.rows:
    if 'Add Multiple' in row.cells[0].text:
        am_row = row
        break

if am_row:
    set_cell_content(am_row.cells[1], [
        ('Opens a combined interaction with two modes — both must be available:', False),
        ('', False),
        ('Mode 1 — Multi-Select Item Lookup (browse users):', False),
        ('Opens item search modal. User selects multiple items using checkboxes. Each selected item loads as a separate individual line row in the grid.', False),
        ('', False),
        ('Mode 2 — Quick Row-Count Input (experienced/power users):', False),
        ('Prompts for a number (N). Inserts N blank rows. Cursor lands on Item Code field of the first new row. User tabs through to fill each row.', False),
        ('', False),
        ('Named Customer Requirement: Pallvaa Mills, SKS', False),
        ('(Requirement originated by Sasi in FSD review — Finding #15 and #16, R2.0)', False),
    ])
    print('FIX 5 done: Add Multiple row expanded')
else:
    print('FIX 5 WARNING: Add Multiple row not found')

# ============================================================
# FIX 6: Replace Section 5.9 content
# ============================================================
para_59 = None
for p in doc.paragraphs:
    if '5.9' in p.text and 'PR List Screen' in p.text and 'Companion' in p.text:
        para_59 = p
        break

if para_59:
    set_para(para_59, '5.9 PR List Screen — Companion Document Register')
    insert_after(para_59._element, [
        ('', False),
        ('CEO Decision (07 May 2026 — Finding #18): The PR List Screen is a new FSD screen requirement accepted by CEO as part of the PR module scope.', False),
        ('', False),
        ('COMPANION SCREEN: PR List Screen', False),
        ('', False),
        ('A document register screen is required for the Purchase Requisition module. This screen displays all PR documents with the following confirmed features:', False),
        ('  • Approval Status badge column (colour-coded per Section 10 badge standards)', False),
        ('  • Filter by status', False),
        ('  • Pagination', False),
        ('', False),
        ('Pending items (not yet included in this Blueprint version):', False),
        ('  • Full column list — Seenivasan and Palanivel’s input due Fri 8 May 2026 to Mariyaiya. Formal column specification to be added in Blueprint v1.2 after FSD amendment is approved and IST column requirements confirmed.', False),
        ('', False),
        ('Timeline: List screen specification will be formally included in Blueprint v1.2, post Sasi review and CEO final approval of this version.', False),
    ])
    print('FIX 6 done: Section 5.9 content expanded')
else:
    print('FIX 6 WARNING: Section 5.9 para not found')

# ============================================================
# FIX 7: Add PR Type Default as new row to Table[40]
# ============================================================
tbl40 = doc.tables[40]
last_row = tbl40.rows[-1]
new_tr = OxmlElement('w:tr')
for ci, src_cell in enumerate(last_row.cells):
    new_tc = OxmlElement('w:tc')
    src_tcPr = src_cell._tc.find(qn('w:tcPr'))
    if src_tcPr is not None:
        new_tc.append(deepcopy(src_tcPr))
    col_text = (
        'PR Type Default' if ci == 0 else
        'Customer / division-configurable at implementation time. Not hardcoded in the application. '
        'Set per customer profile during implementation setup. Default value varies by customer and division.'
    )
    new_tc.append(make_para(col_text))
    new_tr.append(new_tc)
tbl40._tbl.append(new_tr)
print('FIX 7 done: PR Type Default row added to Section 14 table')

# ============================================================
# FIX 8: Update footer Version 1.0 -> 1.1
# ============================================================
updated_footer = False
for section in doc.sections:
    footer = section.footer
    for para in footer.paragraphs:
        for run in para.runs:
            if 'Version 1.0' in run.text:
                run.text = run.text.replace('Version 1.0', 'Version 1.1')
                updated_footer = True
        # Also check if it's in a table in footer
    for tbl in footer.tables:
        for row in tbl.rows:
            for cell in row.cells:
                for para in cell.paragraphs:
                    for run in para.runs:
                        if 'Version 1.0' in run.text:
                            run.text = run.text.replace('Version 1.0', 'Version 1.1')
                            updated_footer = True
if updated_footer:
    print('FIX 8 done: Footer Version updated to 1.1')
else:
    print('FIX 8 WARNING: Footer version not found in runs, trying XML approach')
    for section in doc.sections:
        footer = section.footer
        footer_xml = footer._element
        for t_elem in footer_xml.iter(qn('w:t')):
            if t_elem.text and 'Version 1.0' in t_elem.text:
                t_elem.text = t_elem.text.replace('Version 1.0', 'Version 1.1')
                print('FIX 8 done via XML: Footer updated')

# ============================================================
# FIX 9: Update Document End line v1.0 -> v1.1
# ============================================================
for p in doc.paragraphs:
    if 'Document End' in p.text and 'v1.0' in p.text:
        for run in p.runs:
            if 'v1.0' in run.text:
                run.text = run.text.replace('v1.0', 'v1.1')
        # Also try XML approach
        for t_elem in p._element.iter(qn('w:t')):
            if t_elem.text and 'v1.0' in t_elem.text:
                t_elem.text = t_elem.text.replace('v1.0', 'v1.1')
        print('FIX 9 done: Document End line updated to v1.1')
        break

# ============================================================
# SAVE
# ============================================================
doc.save(r'D:\Spinrise\Docs\UI_Blueprint\SPINRISE_UIDesign_Blueprint_v1.1_FINAL_08May2026.docx')
print()
print('ALL FIXES APPLIED. DOCUMENT SAVED.')
